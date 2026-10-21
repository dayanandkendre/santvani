#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Compare Dnyaneshwari Adhyay 06 and 07 and export an Excel workbook.
Run this script from your website/project root folder.

Expected files:
  pages/granth/dnyaneshwari/adhyay-06.html
  pages/granth/dnyaneshwari/adhyay-07.html

Output:
  dnyaneshwari_adhyay06_07_comparison.xlsx

This script only reads files and creates a report. It does NOT edit HTML or ovi text.
Install dependency once:
  python -m pip install openpyxl
Run:
  python compare_dnyaneshwari_06_07.py
Or specify project root:
  python compare_dnyaneshwari_06_07.py --root "C:\\path\\to\\your\\website"
"""

import argparse
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from html.parser import HTMLParser

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl install nahi. Run: python -m pip install openpyxl")
    sys.exit(1)

SKIP_TAGS = {
    "script", "style", "noscript", "nav", "footer", "header", "form",
    "button", "select", "svg", "iframe", "head", "title"
}
BLOCK_TAGS = {
    "p", "div", "li", "ul", "ol", "br", "h1", "h2", "h3", "h4", "h5", "h6",
    "section", "article", "tr", "td", "th", "blockquote", "pre", "table",
    "main", "aside", "hr", "dd", "dt"
}
VOID_TAGS = {"br", "hr", "img", "input", "meta", "link", "area", "base",
             "embed", "source", "track", "wbr"}
ZERO_WIDTH = dict.fromkeys(map(ord, "\u200b\u200c\u200d\ufeff\u00a0"), " ")


class VisibleTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.skip_stack = []
        self.buffer = []
        self.blocks = []

    def flush(self):
        value = " ".join("".join(self.buffer).split())
        if value:
            self.blocks.append(value)
        self.buffer = []

    def handle_starttag(self, tag, attrs):
        if self.skip_stack:
            if tag in SKIP_TAGS and tag not in VOID_TAGS:
                self.skip_stack.append(tag)
            return
        if tag in SKIP_TAGS and tag not in VOID_TAGS:
            self.flush()
            self.skip_stack.append(tag)
            return
        if tag in BLOCK_TAGS:
            self.flush()

    def handle_startendtag(self, tag, attrs):
        if not self.skip_stack and tag in BLOCK_TAGS:
            self.flush()

    def handle_endtag(self, tag):
        if self.skip_stack:
            if tag in self.skip_stack:
                idx = len(self.skip_stack) - 1 - self.skip_stack[::-1].index(tag)
                self.skip_stack = self.skip_stack[:idx]
                if not self.skip_stack:
                    self.flush()
            return
        if tag in BLOCK_TAGS:
            self.flush()

    def handle_data(self, data):
        if not self.skip_stack:
            self.buffer.append(data)

    def close(self):
        super().close()
        self.flush()


def extract_blocks(path):
    with open(path, "rb") as f:
        raw = f.read()
    text = None
    for enc in ("utf-8-sig", "utf-8"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")

    parser = VisibleTextExtractor()
    parser.feed(text)
    parser.close()
    return parser.blocks


def normalize(text):
    # Conservative: normalize Unicode and whitespace only; do not remove punctuation/digits.
    text = unicodedata.normalize("NFC", text).translate(ZERO_WIDTH).casefold()
    return " ".join(text.split())


def safe_sheet_title(title):
    return title[:31]


def add_table_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(safe_sheet_title(title))
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_sheet(ws)
    return ws


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7A1F2B")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        for row_idx in range(1, min(ws.max_row, 250) + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                max_len = max(max_len, min(len(str(val)), 70))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 14), 55)
    # Long source text columns should be comfortably wide.
    for idx in range(1, ws.max_column + 1):
        header = str(ws.cell(1, idx).value or "").lower()
        if "text" in header or "block" in header or " मजकूर" in header or "ओवी" in header:
            ws.column_dimensions[get_column_letter(idx)].width = 65


def main():
    ap = argparse.ArgumentParser(description="Compare Dnyaneshwari chapters 6 and 7 into Excel.")
    ap.add_argument("--root", default=".", help="Website/project root")
    ap.add_argument("--out", default="dnyaneshwari_adhyay06_07_comparison.xlsx")
    args = ap.parse_args()

    root = os.path.abspath(args.root)
    file6 = os.path.join(root, "pages", "granth", "dnyaneshwari", "adhyay-06.html")
    file7 = os.path.join(root, "pages", "granth", "dnyaneshwari", "adhyay-07.html")

    missing = [p for p in (file6, file7) if not os.path.isfile(p)]
    if missing:
        print("Ya file sapadlya nahit. Script website/project root folder madhun run kara:")
        for path in missing:
            print("  " + path)
        print("\nFolder layout verify kara, kinva --root ne project root dya.")
        sys.exit(2)

    blocks6 = extract_blocks(file6)
    blocks7 = extract_blocks(file7)
    norm6 = [normalize(b) for b in blocks6]
    norm7 = [normalize(b) for b in blocks7]

    counts6, counts7 = Counter(norm6), Counter(norm7)
    text6, text7 = {}, {}
    for n, b in zip(norm6, blocks6):
        text6.setdefault(n, b)
    for n, b in zip(norm7, blocks7):
        text7.setdefault(n, b)

    common_keys = set(counts6) & set(counts7)
    only6_keys = set(counts6) - set(counts7)
    only7_keys = set(counts7) - set(counts6)

    # Include block occurrence details and a simple status for every extracted block.
    rows6 = []
    for i, block in enumerate(blocks6, start=1):
        n = norm6[i - 1]
        status = "Also in Adhyay 07" if n in counts7 else "Only in Adhyay 06"
        rows6.append([i, status, counts6[n], counts7.get(n, 0), block])

    rows7 = []
    for i, block in enumerate(blocks7, start=1):
        n = norm7[i - 1]
        status = "Also in Adhyay 06" if n in counts6 else "Only in Adhyay 07"
        rows7.append([i, status, counts7[n], counts6.get(n, 0), block])

    common_rows = []
    for n in sorted(common_keys, key=lambda x: (text6.get(x, x))):
        common_rows.append([
            min(counts6[n], counts7[n]),
            counts6[n],
            counts7[n],
            "Exact normalized text match",
            text6[n],
            text7[n]
        ])

    only6_rows = [[counts6[n], text6[n]] for n in sorted(only6_keys, key=lambda x: text6[x])]
    only7_rows = [[counts7[n], text7[n]] for n in sorted(only7_keys, key=lambda x: text7[x])]

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Item", "Result"])
    summary_rows = [
        ["Chapter 6 file", os.path.relpath(file6, root).replace("\\", "/")],
        ["Chapter 7 file", os.path.relpath(file7, root).replace("\\", "/")],
        ["Extracted text blocks — Chapter 6", len(blocks6)],
        ["Extracted text blocks — Chapter 7", len(blocks7)],
        ["Unique normalized blocks — Chapter 6", len(counts6)],
        ["Unique normalized blocks — Chapter 7", len(counts7)],
        ["Unique blocks found in both chapters", len(common_keys)],
        ["Unique blocks only in Chapter 6", len(only6_keys)],
        ["Unique blocks only in Chapter 7", len(only7_keys)],
        ["Important note", "A match is a text-level candidate, not proof that an ovi is duplicated. Inspect context and verify against the original granth."],
        ["Safety", "This script reads the two HTML files only. It does not edit source files or correct any text."],
    ]
    for row in summary_rows:
        summary.append(row)
    style_sheet(summary)
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 100

    add_table_sheet(wb, "Adhyay 06 Blocks",
                    ["Block No.", "Comparison status", "Count in Ch. 06", "Count in Ch. 07", "Extracted text"],
                    rows6)
    add_table_sheet(wb, "Adhyay 07 Blocks",
                    ["Block No.", "Comparison status", "Count in Ch. 07", "Count in Ch. 06", "Extracted text"],
                    rows7)
    add_table_sheet(wb, "Common Blocks",
                    ["Shared count", "Count in Ch. 06", "Count in Ch. 07", "Match type", "Chapter 6 text", "Chapter 7 text"],
                    common_rows)
    add_table_sheet(wb, "Only Chapter 06", ["Count in Chapter 06", "Extracted text"], only6_rows)
    add_table_sheet(wb, "Only Chapter 07", ["Count in Chapter 07", "Extracted text"], only7_rows)

    out_path = os.path.abspath(os.path.join(root, args.out))
    wb.save(out_path)
    print("Excel report created:")
    print(out_path)
    print(f"Blocks: chapter 6={len(blocks6)}, chapter 7={len(blocks7)}, common unique blocks={len(common_keys)}")
    print("Note: Excel is for review only; no HTML/ovi content was changed.")


if __name__ == "__main__":
    main()
